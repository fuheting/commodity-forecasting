"""Read and validate configured monthly targets from Excel workbooks."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from .target import (
    InvalidTimestampError,
    MonthlyTargetContract,
    TargetDataError,
    TargetObservation,
    TargetRow,
    TargetSelectionError,
    serialize_numeric,
    validate_rows,
)

MONTH_TOKEN = re.compile(r"^(\d{4})M(0[1-9]|1[0-2])$")
MONTH_LIKE_TOKEN = re.compile(r"^\d{4}M")


@dataclass(frozen=True)
class WorkbookTargetSource:
    """Explicit workbook layout and immutable-source expectations."""

    worksheet_name: str
    target_column: str
    expected_sha256: str
    target: MonthlyTargetContract
    period_column_index: int = 1


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_month_token(value: object) -> date:
    if not isinstance(value, str):
        raise InvalidTimestampError("monthly timestamp must be a string in YYYYMmm form")
    match = MONTH_TOKEN.fullmatch(value)
    if match is None:
        raise InvalidTimestampError(f"invalid monthly timestamp: {value!r}")
    return date(int(match.group(1)), int(match.group(2)), 1)


def _find_target_header(worksheet: Any, target_column: str) -> tuple[int, int]:
    positions: list[tuple[int, int]] = []
    for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        for column_index, value in enumerate(row, start=1):
            if value == target_column:
                positions.append((row_index, column_index))
    if not positions:
        raise TargetSelectionError(f"target column not found: {target_column}")
    if len(positions) != 1:
        raise TargetSelectionError(
            f"target column must appear exactly once, found {len(positions)}"
        )
    return positions[0]


def _read_target_rows(
    worksheet: Any,
    *,
    source: WorkbookTargetSource,
    header_row: int,
    target_column_index: int,
) -> tuple[TargetRow, ...]:
    rows: list[TargetRow] = []
    periods_started = False
    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        period_offset = source.period_column_index - 1
        period_value = row[period_offset] if len(row) > period_offset else None
        if isinstance(period_value, str) and MONTH_LIKE_TOKEN.match(period_value):
            period = parse_month_token(period_value)
            periods_started = True
            target_value = row[target_column_index - 1]
            rows.append(
                TargetRow(source.target.unique_id, period, serialize_numeric(target_value))
            )
        elif periods_started and period_value not in {None, ""}:
            raise InvalidTimestampError(
                f"invalid monthly timestamp after data began: {period_value!r}"
            )
    if not rows:
        raise InvalidTimestampError("workbook contains no monthly timestamp rows")
    return tuple(rows)


def extract_workbook_target(
    workbook_path: Path,
    source: WorkbookTargetSource,
) -> TargetObservation:
    """Read the configured monthly target without mutating its source workbook."""

    if not workbook_path.is_file():
        raise TargetDataError(f"raw workbook does not exist: {workbook_path}")
    source_sha256_before = sha256_file(workbook_path)
    if source_sha256_before != source.expected_sha256:
        raise TargetDataError("raw workbook hash does not match the configured source")

    workbook: Any | None = None
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if source.worksheet_name not in workbook.sheetnames:
            raise TargetSelectionError(f"worksheet not found: {source.worksheet_name}")
        worksheet = workbook[source.worksheet_name]
        header_row, target_column_index = _find_target_header(
            worksheet, source.target_column
        )
        rows = _read_target_rows(
            worksheet,
            source=source,
            header_row=header_row,
            target_column_index=target_column_index,
        )
    finally:
        if workbook is not None:
            workbook.close()

    source_sha256_after = sha256_file(workbook_path)
    if source_sha256_after != source_sha256_before:
        raise TargetDataError("raw workbook hash changed during target extraction")
    validate_rows(rows, source.target)
    first, last = rows[0].ds, rows[-1].ds
    return TargetObservation(
        rows=rows,
        header_row=header_row,
        target_column_index=target_column_index,
        period_start=f"{first.year:04d}M{first.month:02d}",
        period_end=f"{last.year:04d}M{last.month:02d}",
        source_sha256_before=source_sha256_before,
        source_sha256_after=source_sha256_after,
    )
